# ==========================================================
# Genera data/geovisor_datos.xlsx y data/municipios.geojson
# a partir del MGN 2018 del DANE (municipios).
# Los valores de recurso son ILUSTRATIVOS: reemplazar con NASA POWER.
# Uso: python scripts/construir_datos.py ruta/co_2018_MGN_MPIO_POLITICO.geojson
# ==========================================================
import json
import math
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from shapely.geometry import mapping, shape

random.seed(7)
RAIZ = Path(__file__).resolve().parents[1]
MGN = sys.argv[1] if len(sys.argv) > 1 else "co_2018_MGN_MPIO_POLITICO.geojson"
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

# ---- Nombres legibles
ESPECIALES = {"ARCHIPIÉLAGO DE SAN ANDRÉS, PROVIDENCIA Y SANTA CATALINA": "San Andrés y Providencia",
              "BOGOTÁ, D.C.": "Bogotá, D.C.", "QUINDIO": "Quindío"}


def legible(n):
    if n in ESPECIALES:
        return ESPECIALES[n]
    s = n.title()
    for w in [" De ", " Del ", " La ", " El ", " Los ", " Las ", " Y "]:
        s = s.replace(w, w.lower())
    return s


# ---- Región natural (aproximación; validar con la lista oficial del IGAC)
CARIBE = {"08", "13", "20", "23", "44", "47", "70"}
ORINOQUIA = {"81", "85", "50", "99"}
AMAZONIA = {"91", "18", "94", "95", "86", "97"}


def region(dp, lon, lat):
    if dp == "88":
        return "Insular"
    if dp in CARIBE:
        return "Caribe"
    if dp in ORINOQUIA:
        return "Orinoquía"
    if dp in AMAZONIA:
        return "Amazonía"
    if dp == "27":
        return "Pacífica"
    if dp == "05":
        return "Caribe" if lat >= 7.6 else "Andina"
    if dp == "76":
        return "Andina" if lon > -76.75 else "Pacífica"
    if dp == "19":
        return "Andina" if lon > -77.0 else "Pacífica"
    if dp == "52":
        return "Andina" if (lon > -77.95 and lat < 1.9) else "Pacífica"
    return "Andina"


# ---- Centros volcánicos de referencia (coordenadas aproximadas)
VOLCANES = [("Nevado del Ruiz", 4.892, -75.324), ("Nevado de Santa Isabel", 4.82, -75.37),
            ("Nevado del Tolima", 4.66, -75.33), ("Cerro Machín", 4.49, -75.39),
            ("Nevado del Huila", 2.93, -76.03), ("Puracé", 2.32, -76.40), ("Doña Juana", 1.50, -76.94),
            ("Galeras", 1.22, -77.36), ("Azufral", 1.09, -77.72), ("Cumbal", 0.95, -77.88),
            ("Chiles", 0.82, -77.93), ("Paipa-Iza", 5.78, -73.12)]


def km(a, b, c, d):
    p1, p2 = math.radians(a), math.radians(c)
    h = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(d - b) / 2) ** 2
    return 2 * 6371 * math.asin(math.sqrt(h))


# ---- Recurso ilustrativo (estructura espacial plausible, NO medida)
def ghi_anual(lon, lat):
    v = 4.3 + 1.5 / (1 + math.exp(-(lat - 8.3)))
    if lon > -73.6 and lat > 10.4:
        v += 0.45
    if lon > -72.5 and 3 < lat < 7.2:
        v += 0.35
    if lon < -76.6 and lat < 8:
        v -= 0.6
    if lat < 1.2 and lon > -74:
        v -= 0.15
    return min(max(v + random.gauss(0, 0.08), 3.5), 6.6)


def viento_anual(lon, lat, reg):
    if lon > -73.4 and lat > 11.0:
        v = 7.2 + 1.2 * min((lat - 11) / 0.9, 1)
    elif reg == "Caribe":
        v = 4.0 + 0.8 * max(0, (lat - 9) / 2)
    elif reg == "Orinoquía":
        v = 3.8
    elif reg == "Amazonía":
        v = 2.0
    elif reg == "Pacífica":
        v = 2.3
    elif reg == "Insular":
        v = 6.5
    else:
        v = 2.9
    return max(v + random.gauss(0, 0.25), 1.2)


PATRON_GHI = {
    "Caribe":    [1.05, 1.08, 1.10, 1.03, 0.97, 0.98, 1.02, 1.01, 0.95, 0.90, 0.92, 0.99],
    "Andina":    [1.04, 1.03, 0.98, 0.93, 0.94, 1.00, 1.07, 1.06, 1.01, 0.94, 0.93, 0.98],
    "Orinoquía": [1.15, 1.15, 1.08, 0.97, 0.92, 0.88, 0.90, 0.94, 0.98, 0.99, 1.00, 1.06],
    "Amazonía":  [1.00, 0.98, 0.97, 0.96, 0.98, 1.00, 1.04, 1.07, 1.05, 1.01, 0.98, 0.97],
    "Pacífica":  [0.98, 1.00, 1.02, 1.01, 0.98, 0.97, 1.02, 1.02, 1.00, 0.98, 0.99, 1.01],
}
PATRON_VIENTO = {
    "Caribe":    [1.12, 1.15, 1.10, 0.98, 0.92, 1.08, 1.18, 1.10, 0.88, 0.72, 0.75, 0.95],
    "Andina":    [0.98, 0.98, 0.97, 0.95, 0.97, 1.08, 1.12, 1.10, 1.02, 0.94, 0.93, 0.96],
    "Orinoquía": [1.25, 1.25, 1.15, 0.95, 0.85, 0.85, 0.88, 0.88, 0.90, 0.95, 1.02, 1.12],
    "Amazonía":  [1.02, 1.02, 1.00, 0.98, 0.97, 0.98, 1.00, 1.02, 1.02, 1.00, 0.99, 1.00],
    "Pacífica":  [1.00, 1.02, 1.02, 1.00, 0.98, 0.98, 1.00, 1.00, 1.00, 0.99, 1.00, 1.01],
}
PATRON_GHI["Insular"] = PATRON_GHI["Caribe"]
PATRON_VIENTO["Insular"] = PATRON_VIENTO["Caribe"]


def mensual(anual, patron, ruido):
    f = [p * (1 + random.gauss(0, ruido)) for p in patron]
    m = sum(f) / 12
    return [round(anual * x / m, 2) for x in f]      # el promedio de los 12 meses = valor anual


# ---- Construcción de filas
mgn = json.load(open(MGN, encoding="utf-8"))
filas, geo_feats = [], []
for f in mgn["features"]:
    p = f["properties"]
    sh = shape(f["geometry"])
    c = sh.representative_point()
    lon, lat = c.x, c.y
    cod, dp = p["MPIO_CCNCT"], p["DPTO_CCDGO"]
    reg = region(dp, lon, lat)
    nombre_v, la_v, lo_v = min(VOLCANES, key=lambda v: km(lat, lon, v[1], v[2]))
    dist = km(lat, lon, la_v, lo_v)
    andina = reg == "Andina"
    filas.append({
        "cod_dane": cod,
        "departamento": legible(p["DPTO_CNMBR"]), "municipio": legible(p["MPIO_CNMBR"]),
        "latitud": round(lat, 5), "longitud": round(lon, 5), "region_natural": reg,
        "densidad_aire": 1.00 if andina else 1.17,
        "weibull_k": 2.8 if (lon > -73.4 and lat > 11.0) else 2.0,
        "indice_geotermico": round(math.exp(-dist / 60), 3) if andina else None,
        "volcan_cercano": nombre_v if andina else None,
        "dist_volcan_km": round(dist, 1) if andina else None,
        "ghi_m": mensual(ghi_anual(lon, lat), PATRON_GHI[reg], 0.015),
        "viento_m": mensual(viento_anual(lon, lat, reg), PATRON_VIENTO[reg], 0.02),
    })
    geo = sh.simplify(0.003, preserve_topology=True)
    geo_feats.append({"type": "Feature", "geometry": mapping(geo), "properties": {"cod": cod, "dpto": dp}})
filas.sort(key=lambda r: r["cod_dane"])


# ---- GeoJSON simplificado (coordenadas a 4 decimales, unos 11 m)
def redondear(o):
    if isinstance(o, float):
        return round(o, 4)
    if isinstance(o, (list, tuple)):
        return [redondear(x) for x in o]
    if isinstance(o, dict):
        return {k: redondear(v) for k, v in o.items()}
    return o


(RAIZ / "data").mkdir(exist_ok=True)
with open(RAIZ / "data" / "municipios.geojson", "w", encoding="utf-8") as fh:
    json.dump(redondear({"type": "FeatureCollection", "features": geo_feats}), fh,
              ensure_ascii=False, separators=(",", ":"))

# ---- Excel
AZUL, NEGRO = "0000FF", "000000"
CAB = PatternFill("solid", fgColor="1F3A5F")
AMARILLO = PatternFill("solid", fgColor="FFFF00")


def f_txt(color=NEGRO, **k):
    return Font(name="Arial", size=10, color=color, **k)


def encabezado(ws, cols, anchos):
    ws.append(cols)
    for i, w in enumerate(anchos, 1):
        c = ws.cell(1, i)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


wb = Workbook()
n = len(filas)
ult = n + 1

# Hoja Municipios
ws = wb.active
ws.title = "Municipios"
cols = ["cod_dane", "departamento", "municipio", "latitud", "longitud", "region_natural",
        "zona_central", "ghi_kwh_m2_dia", "viento_50m_m_s", "densidad_aire_kg_m3", "weibull_k",
        "indice_geotermico", "volcan_cercano", "dist_volcan_km", "fuente"]
encabezado(ws, cols, [10, 24, 26, 10, 11, 13, 11, 13, 13, 13, 10, 13, 22, 13, 44])
for r, d in enumerate(filas, 2):
    valores = [d["cod_dane"], d["departamento"], d["municipio"], d["latitud"], d["longitud"],
               d["region_natural"], None, None, None, d["densidad_aire"], d["weibull_k"],
               d["indice_geotermico"], d["volcan_cercano"], d["dist_volcan_km"],
               "Ilustrativo v1 — reemplazar con NASA POWER"]
    for col, v in enumerate(valores, 1):
        ws.cell(r, col, v).font = f_txt(AZUL)
    # Columnas calculadas (negro): zona central y promedios anuales desde las hojas mensuales
    ws.cell(r, 7, f'=IF(F{r}="Andina","Sí","No")').font = f_txt()
    ws.cell(r, 8, f"=ROUND(AVERAGE(INDEX(GHI_mensual!$B$2:$M${ult},MATCH(A{r},GHI_mensual!$A$2:$A${ult},0),0)),2)").font = f_txt()
    ws.cell(r, 9, f"=ROUND(AVERAGE(INDEX(Viento_mensual!$B$2:$M${ult},MATCH(A{r},Viento_mensual!$A$2:$A${ult},0),0)),2)").font = f_txt()
    for col, fmt in [(4, "0.00000"), (5, "0.00000"), (8, "0.00"), (9, "0.00"), (10, "0.00"),
                     (11, "0.0"), (12, "0.000"), (14, "0.0")]:
        ws.cell(r, col).number_format = fmt
ws.auto_filter.ref = f"A1:O{ult}"

# Hojas mensuales
for nombre, clave, unidad in [("GHI_mensual", "ghi_m", "kWh/m²·día"), ("Viento_mensual", "viento_m", "m/s a 50 m")]:
    h = wb.create_sheet(nombre)
    encabezado(h, ["cod_dane"] + [f"{m} ({unidad})" for m in MESES], [10] + [12] * 12)
    for r, d in enumerate(filas, 2):
        h.cell(r, 1, d["cod_dane"]).font = f_txt(AZUL)
        for j, v in enumerate(d[clave], 2):
            c = h.cell(r, j, v)
            c.font = f_txt(AZUL)
            c.number_format = "0.00"

# Tecnologías (costos editables)
t = wb.create_sheet("Tecnologias")
encabezado(t, ["tecnologia_id", "recurso", "nombre", "capex", "unidad_capex", "opex_anual",
               "unidad_opex", "vida_util_anios", "nota"], [16, 11, 30, 10, 12, 11, 16, 12, 70])
TEC = [
    ("fv", "Solar", "Paneles fotovoltaicos", 900, "USD/kWp", 13.5, "USD/kWp·año", 25,
     "Sistema conectado a red, sin almacenamiento"),
    ("colector", "Solar", "Colector solar térmico", 350, "USD/m²", 5, "USD/m²·año", 20,
     "Colector plano para agua caliente o procesos"),
    ("eolico_grande", "Viento", "Aerogenerador gran escala", 1500, "USD/kW", 40, "USD/kW·año", 20,
     "Turbinas de 1 a 5 MW"),
    ("eolico_pequeno", "Viento", "Aerogenerador pequeño", 3500, "USD/kW", 60, "USD/kW·año", 20,
     "Turbinas de 1 a 100 kW para autoconsumo"),
    ("geo_binaria", "Geotermia", "Planta binaria (ORC)", 4500, "USD/kW", 120, "USD/kW·año", 30,
     "Recursos de 100 a 180 °C; incluye perforación"),
    ("geo_flash", "Geotermia", "Planta flash", 3800, "USD/kW", 110, "USD/kW·año", 30,
     "Recursos sobre 180 °C; incluye perforación"),
]
for r, row in enumerate(TEC, 2):
    for col, v in enumerate(row, 1):
        c = t.cell(r, col, v)
        c.font = f_txt(AZUL)
        if col in (4, 6):
            c.fill = AMARILLO
t.cell(len(TEC) + 3, 1, "Costos de referencia (orden de magnitud). Validar con IRENA, UPME o cotizaciones "
       "locales antes de usarlos en decisiones.").font = f_txt(italic=True)

# Supuestos
s = wb.create_sheet("Supuestos")
encabezado(s, ["clave", "parametro", "valor", "unidad", "nota"], [18, 36, 10, 12, 80])
SUP = [
    ("tasa_descuento", "Tasa de descuento", 0.10, "fracción", "Para el factor de recuperación de capital del LCOE"),
    ("pr_fv", "Performance ratio FV", 0.80, "fracción", "Pérdidas por inversor, cableado, temperatura y suciedad"),
    ("alfa", "Exponente de cizalladura (α)", 0.14, "—",
     "Ley de potencia para llevar el viento de 50 m al buje; 0,10 mar, 0,14 abierto, 0,25 rugoso"),
    ("cp", "Coeficiente de potencia (Cp)", 0.45, "fracción", "Eficiencia aerodinámica de la turbina; límite de Betz 0,593"),
    ("perdidas_eolicas", "Pérdidas eólicas", 0.12, "fracción", "Estelas, disponibilidad y eléctricas"),
    ("v_arranque", "Velocidad de arranque", 3.0, "m/s", "Bajo esta velocidad la turbina no produce"),
    ("v_corte", "Velocidad de corte", 25.0, "m/s", "Sobre esta velocidad la turbina se detiene por seguridad"),
    ("eta0_colector", "Eficiencia óptica del colector (η0)", 0.75, "fracción", "Colector plano típico"),
    ("a1_colector", "Coeficiente de pérdidas térmicas (a1)", 3.5, "W/m²·K", "Colector plano típico"),
    ("t_ambiente", "Temperatura ambiente media", 22.0, "°C", "Valor nacional único; reemplazar por dato municipal"),
    ("horas_sol", "Horas de sol por día", 12.0, "h", "Cerca del ecuador el día dura unas 12 h todo el año"),
    ("umbral_geo_alto", "Umbral índice geotérmico Alto", 0.6, "—", "Índice = exp(−d/60 km) a un centro volcánico"),
    ("umbral_geo_medio", "Umbral índice geotérmico Medio", 0.3, "—", "Por debajo se clasifica Bajo"),
]
for r, row in enumerate(SUP, 2):
    for col, v in enumerate(row, 1):
        c = s.cell(r, col, v)
        c.font = f_txt(AZUL if col == 3 else NEGRO)
        if col == 3:
            c.fill = AMARILLO

# Diccionario
d = wb.create_sheet("Diccionario")
encabezado(d, ["hoja / columna", "descripción", "fuente", "tipo"], [30, 80, 44, 12])
DIC = [
    ("Leyenda", "Azul: dato editable. Negro: fórmula. Amarillo: supuesto clave.", "—", "—"),
    ("Municipios / cod_dane", "Código DANE de 5 dígitos; llave para unir con las hojas mensuales", "DANE MGN 2018", "Dato"),
    ("Municipios / latitud, longitud", "Punto interior del polígono municipal (WGS84); ubicación del pin",
     "Calculado del MGN 2018", "Dato"),
    ("Municipios / region_natural", "Región natural; Antioquia, Valle, Cauca y Nariño divididos por posición",
     "Aproximación; validar con IGAC", "Supuesto"),
    ("Municipios / zona_central", "Sí si la región es Andina; define dónde aplica geotermia", "Fórmula", "Calculado"),
    ("Municipios / ghi_kwh_m2_dia", "Promedio de los 12 meses de GHI_mensual", "Fórmula", "Calculado"),
    ("Municipios / viento_50m_m_s", "Promedio de los 12 meses de Viento_mensual", "Fórmula", "Calculado"),
    ("Municipios / densidad_aire_kg_m3", "1,00 en zona andina (unos 1.800 m) y 1,17 en tierras bajas cálidas",
     "Supuesto por región", "Supuesto"),
    ("Municipios / weibull_k", "Forma de la distribución del viento; 2,0 general, 2,8 en alisios de La Guajira",
     "Supuesto", "Supuesto"),
    ("Municipios / indice_geotermico", "exp(−d/60) con d = distancia al centro volcánico más cercano",
     "Supuesto de trabajo", "Estimado"),
    ("GHI_mensual", "Irradiación global horizontal media diaria por mes",
     "Ilustrativo → NASA POWER ALLSKY_SFC_SW_DWN", "Estimado"),
    ("Viento_mensual", "Velocidad media del viento a 50 m por mes", "Ilustrativo → NASA POWER WS50M", "Estimado"),
    ("Tecnologias", "CAPEX, OPEX y vida útil por tecnología; el dashboard los lee de aquí", "Referencia, validar", "Supuesto"),
    ("Supuestos", "Parámetros técnicos y financieros globales; el dashboard los lee por clave",
     "Referencia, validar", "Supuesto"),
]
for row in DIC:
    d.append(list(row))
for row in d.iter_rows(min_row=2):
    for c in row:
        c.font = f_txt()
        c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(RAIZ / "data" / "geovisor_datos.xlsx")
print(f"{n} municipios escritos")
